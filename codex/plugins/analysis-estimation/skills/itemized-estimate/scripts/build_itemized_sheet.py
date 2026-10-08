"""Build a customer-facing itemized estimate sheet on top of a reference .xlsx.

  python build_itemized_sheet.py inspect <template.xlsx>
      Print the template as JSON: sheets, header rows, data rows, notes, formulas.
  python build_itemized_sheet.py build <spec.json>
      Write the sheet described by spec.json.

spec.json:
{
  "template": "path/to/reference.xlsx",
  "sheet": "評估表 ",            # optional; default = first sheet
  "out": "path/to/output.xlsx",
  "rate": 8100,                  # optional; replaces the number in the template's rate formulas
  "items": [                     # already in delivery order; numbered 1..n here
    {"name": "...", "days": 646, "lines": ["1. ...", "2. ...", {"red": "※..."}]}
  ],
  "notes": ["※...", {"red": "※..."}]   # rows under the table, column C
}
Relative paths resolve against the spec file's folder.
"""
import copy
import json
import math
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.formula.translate import Translator

RED = 'FFFF0000'
LINE_PT = 15.375  # one wrapped line at 12pt CJK fonts (Excel autofit)


def text_of(v):
    return str(v) if not isinstance(v, CellRichText) else ''.join(
        b if isinstance(b, str) else b.text for b in v)


def first_data_row(ws):
    for r in range(1, ws.max_row + 1):
        if isinstance(ws.cell(r, 1).value, (int, float)):
            return r
    raise SystemExit('template: no row with a number in column A')


def inspect(path):
    wb = openpyxl.load_workbook(path, rich_text=True)
    out = {'sheets': wb.sheetnames}
    ws = wb[wb.sheetnames[0]]
    first = first_data_row(ws)
    rows = []
    for r in range(1, ws.max_row + 1):
        vals = {c.column_letter: text_of(c.value) for c in ws[r] if c.value is not None}
        if not vals:
            continue
        red = [c.column_letter for c in ws[r] if isinstance(c.value, CellRichText) and any(
            not isinstance(b, str) and b.font.color is not None and b.font.color.rgb == RED for b in c.value)]
        kind = 'header' if r < first else ('item' if 'A' in vals else 'note')
        rows.append({'row': r, 'kind': kind, 'cells': vals, 'red_cells': red})
    out['first_sheet'] = ws.title
    out['first_data_row'] = first
    out['freeze'] = ws.freeze_panes
    out['rows'] = rows
    print(json.dumps(out, ensure_ascii=False, indent=1))


def units(s):
    return sum(2.2 if unicodedata.east_asian_width(ch) in 'WF' else 1.1 for ch in s)


def visual_lines(lines, width):
    cap = width * 0.9
    return sum(max(1, math.ceil(units(l) / cap)) for l in lines)


def cell_value(lines, font):
    plain = [x for x in lines if isinstance(x, str)]
    if len(plain) == len(lines):
        return '\n'.join(lines)
    base = InlineFont(rFont=font.name, sz=font.sz)
    red = InlineFont(rFont=font.name, sz=font.sz, color=RED)
    blocks = []
    for i, x in enumerate(lines):
        sep = '\n' if i < len(lines) - 1 else ''
        if isinstance(x, str):
            blocks.append(TextBlock(base, x + sep))
        else:
            blocks.append(TextBlock(red, x['red'] + sep))
    return CellRichText(blocks)


def flat(lines):
    return [x if isinstance(x, str) else x['red'] for x in lines]


def set_rate(formula, rate):
    return re.sub(r'\*\s*\d+(\.\d+)?', f'*{rate}', formula) if rate else formula


def build(spec_path):
    spec_path = Path(spec_path).resolve()
    spec = json.loads(spec_path.read_text(encoding='utf-8'))
    base = spec_path.parent
    tpl = (base / spec['template']).resolve()
    out = (base / spec['out']).resolve()
    rate = spec.get('rate')

    wb = openpyxl.load_workbook(tpl, rich_text=True)
    name = spec.get('sheet') or wb.sheetnames[0]
    for other in [s for s in wb.sheetnames if s != name]:
        del wb[other]
    ws = wb[name]
    ws.title = name.strip()
    first = first_data_row(ws)

    cols = [c.column_letter for c in ws[first]]
    tmpl_style = {c: copy.copy(ws[f'{c}{first}']._style) for c in cols}
    tmpl_formula = {c: ws[f'{c}{first}'].value for c in cols
                    if isinstance(ws[f'{c}{first}'].value, str) and ws[f'{c}{first}'].value.startswith('=')}
    for row in ws.iter_rows(min_row=first, max_row=ws.max_row):
        for c in row:
            c.value = None
    width_c = ws.column_dimensions['C'].width or 60

    r = first
    for no, item in enumerate(spec['items'], 1):
        for c in cols:
            ws[f'{c}{r}']._style = copy.copy(tmpl_style[c])
        ws[f'A{r}'] = no
        ws[f'B{r}'] = item['name']
        ws[f'C{r}'] = cell_value(item['lines'], ws[f'C{r}'].font)
        ws[f'D{r}'] = item['days']
        for c, f in tmpl_formula.items():
            if c not in 'ABCD':
                ws[f'{c}{r}'] = set_rate(Translator(f, origin=f'{c}{first}').translate_formula(f'{c}{r}'), rate)
        ws.row_dimensions[r].height = round(LINE_PT * visual_lines(flat(item["lines"]), width_c) * 20) / 20
        r += 1
    last = r - 1

    for note in spec.get('notes', []):
        ws[f'C{r}']._style = copy.copy(tmpl_style['C'])
        ws[f'C{r}'] = cell_value([note], ws[f'C{r}'].font)
        ws.row_dimensions[r].height = round(LINE_PT * visual_lines(flat([note]), width_c) * 20) / 20
        r += 1

    for hr in range(1, first):
        for c in ws[hr]:
            if isinstance(c.value, str) and c.value.startswith('='):
                v = re.sub(rf'([A-Z]+){first}:([A-Z]+)\d+', rf'\g<1>{first}:\g<2>{last}', c.value)
                c.value = set_rate(v, rate)

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(out, 'items', len(spec['items']), 'total', sum(i['days'] for i in spec['items']))


if __name__ == '__main__':
    cmd, arg = sys.argv[1], sys.argv[2]
    inspect(arg) if cmd == 'inspect' else build(arg)
